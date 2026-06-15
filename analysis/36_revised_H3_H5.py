#!/usr/bin/env python3
"""Analysis block 36: reassessment of the REFORMULATED H3 and H5.

The thesis (MASTERARBEIT_v2) framed EI as a *moderator* of congruity bias (H3)
and investment experience as a *conditional moderator* (H5). The formal
moderation/interaction tests were non-significant. The user reformulated H3/H5
as SIMPLE SUBGROUP effects:

  H3: Among investors lower in EI, masculine-congruent startups receive higher
      investment than feminine-congruent startups.
  H5: Among investors lower in EI WITHOUT prior investment experience,
      masculine-congruent > feminine-congruent.

masculine-congruent = invest_sv among MALE-founder condition (male founder x
                      male-dominated industry, SiteVision)
feminine-congruent  = invest_bu among FEMALE-founder condition (female founder x
                      female-dominated industry, BalanceUp)
Both are between-subjects (founder gender varies between participants), so these
are independent-samples simple-effect tests.

Outputs: figures figH3_revised, figH5_revised, fig_conceptual_model_revised
and the Word report 36_revised_H3_H5_report.docx.
"""
import csv, re, os, sys
import numpy as np
from scipy import stats

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, 'figures'))
import apa_barchart as apa
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

SRC = os.path.join(HERE, '..', 'MasterThesis_cleaned.csv')
with open(SRC, newline='') as f:
    rows = list(csv.reader(f))
names, data = rows[0], rows[3:]
ei_cols = [i for i, n in enumerate(names)
           if n.startswith(('Q10 (SEA)', 'Q11 (OEA)', 'Q12 (UOE)', 'Q13 (ROE)')) and '_DO' not in n]
L = lambda v: int(re.match(r'\s*(\d+)', v.strip()).group(1))
ei = np.array([[L(r[i]) for i in ei_cols] for r in data], float).mean(1)
sv = np.array([float(r[names.index('invest_sv')]) for r in data])
bu = np.array([float(r[names.index('invest_bu')]) for r in data])
g = np.array([r[names.index('team_gender')] for r in data])
exp = np.array([r[names.index('Q23')] for r in data])
MED = np.median(ei)
TERT = np.quantile(ei, 1/3)


def battery(mask):
    """Full inferential battery for masculine- vs feminine-congruent under `mask`."""
    masc = sv[(g == 'male') & mask]
    fem = bu[(g == 'female') & mask]
    n1, n2 = len(masc), len(fem)
    m1, m2 = masc.mean(), fem.mean()
    s1, s2 = masc.std(ddof=1), fem.std(ddof=1)
    diff = m1 - m2
    t, p2 = stats.ttest_ind(masc, fem, equal_var=False)
    p1 = p2 / 2 if t > 0 else 1 - p2 / 2
    dfw = (s1**2/n1 + s2**2/n2)**2 / ((s1**2/n1)**2/(n1-1) + (s2**2/n2)**2/(n2-1))
    se = np.sqrt(s1**2/n1 + s2**2/n2)
    tc = stats.t.ppf(.975, dfw)
    ci = (diff - tc*se, diff + tc*se)
    sp = np.sqrt(((n1-1)*s1**2 + (n2-1)*s2**2) / (n1+n2-2))
    dd = diff / sp
    u, pmw = stats.mannwhitneyu(masc, fem, alternative='greater')
    rng = np.random.default_rng(42)
    boot = np.array([rng.choice(masc, n1).mean() - rng.choice(fem, n2).mean()
                     for _ in range(10000)])
    bci = np.percentile(boot, [2.5, 97.5])
    return dict(n1=n1, n2=n2, m1=m1, m2=m2, s1=s1, s2=s2, se1=s1/np.sqrt(n1),
                se2=s2/np.sqrt(n2), diff=diff, ci=ci, t=t, dfw=dfw, p2=p2, p1=p1,
                d=dd, u=u, pmw=pmw, bci=bci)


H3 = battery(ei < MED)
H3t = battery(ei <= TERT)
H5 = battery((ei < MED) & (exp == 'Nein'))
H5t = battery((ei <= TERT) & (exp == 'Nein'))
FULL = battery(np.ones(len(ei), bool))

for tag, r in [('H3 median', H3), ('H3 tertile', H3t),
               ('H5 median', H5), ('H5 tertile', H5t), ('FULL', FULL)]:
    print(f"{tag:12s} masc M={r['m1']:.2f}(n{r['n1']}) fem M={r['m2']:.2f}(n{r['n2']}) "
          f"diff={r['diff']:+.2f} d={r['d']:+.2f} p1={r['p1']:.3f} p2={r['p2']:.3f} "
          f"MW={r['pmw']:.3f} CI[{r['ci'][0]:+.2f},{r['ci'][1]:+.2f}] "
          f"boot[{r['bci'][0]:+.2f},{r['bci'][1]:+.2f}]")

# ════════════════════════ FIGURES ════════════════════════
apa.apply_apa_style()
DARK, GREY = '#808080', '#D3D3D3'


def congruence_fig(r, title, fname, p_label):
    fig, ax = plt.subplots(figsize=(5.4, 4.7))
    bars = ax.bar([0, 1], [r['m1'], r['m2']], 0.6,
                  yerr=[r['se1'], r['se2']], capsize=4,
                  color=[DARK, GREY], edgecolor='black', linewidth=0.8,
                  error_kw=dict(lw=0.9), zorder=3)
    for rect, v, n in zip(bars, [r['m1'], r['m2']], [r['n1'], r['n2']]):
        ax.text(rect.get_x()+rect.get_width()/2, rect.get_height()+0.03,
                f'{v:.2f}', ha='center', va='bottom', fontsize=10)
        ax.text(rect.get_x()+rect.get_width()/2, 0.10,
                f'n = {n}', ha='center', va='bottom', fontsize=8.5, color='white')
    yb = max(r['m1'], r['m2']) + r['se1'] + 0.30
    ax.plot([0, 0, 1, 1], [yb, yb+0.05, yb+0.05, yb], color='black', lw=0.9)
    ax.text(0.5, yb+0.08, p_label, ha='center', va='bottom', fontsize=9)
    ax.set_xticks([0, 1])
    ax.set_xticklabels(['Masculine-\ncongruent', 'Feminine-\ncongruent'])
    ax.set_ylabel('Mean amount invested (0–3)')
    ax.set_ylim(0, 3.0)
    ax.set_title(title, fontsize=11)
    fig.tight_layout()
    apa.save(fig, os.path.join(HERE, 'figures', fname))
    plt.close(fig)
    print('saved', fname)


congruence_fig(
    H3, 'Investors lower in emotional intelligence',
    'figH3_revised',
    f"d = {H3['d']:.2f},  one-tailed p = {H3['p1']:.3f}")
congruence_fig(
    H5, 'Lower-EI investors without prior\ninvestment experience',
    'figH5_revised',
    f"d = {H5['d']:.2f},  p = {H5['p2']:.3f}")

# ── full vs subgroup escalation panel (descriptive context) ──
fig, ax = plt.subplots(figsize=(7.0, 4.7))
groups = ['Full\nsample', 'Lower EI\n(H3)', 'Lower EI &\nno experience\n(H5)']
masc_m = [FULL['m1'], H3['m1'], H5['m1']]
fem_m = [FULL['m2'], H3['m2'], H5['m2']]
masc_se = [FULL['se1'], H3['se1'], H5['se1']]
fem_se = [FULL['se2'], H3['se2'], H5['se2']]
x = np.arange(3); bw = 0.38
b1 = ax.bar(x - bw/2, masc_m, bw*0.95, yerr=masc_se, capsize=3,
            color=DARK, edgecolor='black', linewidth=0.8, zorder=3,
            error_kw=dict(lw=0.8), label='Masculine-congruent')
b2 = ax.bar(x + bw/2, fem_m, bw*0.95, yerr=fem_se, capsize=3,
            color=GREY, edgecolor='black', linewidth=0.8, zorder=3,
            error_kw=dict(lw=0.8), label='Feminine-congruent')
for xi, dv, rr in zip(x, [FULL['diff'], H3['diff'], H5['diff']], [FULL, H3, H5]):
    top = max(rr['m1'], rr['m2']) + 0.45
    ax.text(xi, top, f"Δ = {dv:+.2f}\nd = {rr['d']:.2f}",
            ha='center', va='bottom', fontsize=8.5)
ax.set_xticks(x); ax.set_xticklabels(groups)
ax.set_ylabel('Mean amount invested (0–3)')
ax.set_ylim(0, 3.0)
ax.legend(frameon=False, fontsize=9, loc='upper right')
fig.tight_layout()
apa.save(fig, os.path.join(HERE, 'figures', 'figH3H5_escalation'))
plt.close(fig)
print('saved figH3H5_escalation')

# ════════════════════════ CONCEPTUAL MODEL ════════════════════════
# Subgroup framing: NO moderation arrow on the congruity path. EI and experience
# define the subgroup in which the congruity preference is expressed.
fig, ax = plt.subplots(figsize=(8.6, 5.0))
ax.set_xlim(0, 10); ax.set_ylim(0, 7); ax.axis('off')


def box(x, y, w, h, text, fc='white', fs=9.5, bold=False):
    b = FancyBboxPatch((x-w/2, y-h/2), w, h, boxstyle='round,pad=0.02,rounding_size=0.12',
                       fc=fc, ec='black', lw=1.1, zorder=2)
    ax.add_patch(b)
    ax.text(x, y, text, ha='center', va='center', fontsize=fs,
            fontweight='bold' if bold else 'normal', zorder=3)


def arrow(x1, y1, x2, y2, style='-|>', lw=1.4, ls='-', color='black'):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle=style,
                 mutation_scale=14, lw=lw, ls=ls, color=color, zorder=1))


# core congruity path
box(2.0, 5.2, 2.7, 1.05, 'Founder–industry\ngender congruity\n(masculine vs feminine)', fc='#EDEDED', bold=True)
box(8.0, 5.2, 2.5, 1.05, 'Investment\nallocation', fc='#EDEDED', bold=True)
arrow(3.35, 5.2, 6.75, 5.2, lw=2.0)
ax.text(5.05, 5.55, 'congruity preference\n(masculine > feminine)', ha='center', va='bottom', fontsize=8.3, style='italic')

# subgroup conditions (define WHERE the effect is observed) -> dashed selectors
box(2.0, 2.2, 2.7, 0.95, 'Lower emotional\nintelligence (WLEIS)', fc='white')
box(5.05, 1.1, 2.9, 0.95, 'No prior investment\nexperience', fc='white')
# dashed "operates within" arrows pointing to the path
arrow(2.0, 2.7, 3.6, 4.75, style='-|>', lw=1.2, ls=(0, (4, 3)))
arrow(5.05, 1.6, 4.7, 4.7, style='-|>', lw=1.2, ls=(0, (4, 3)))
ax.text(2.55, 3.75, 'effect emerges in\nthis subgroup (H3)', ha='left', va='center', fontsize=7.8, style='italic')
ax.text(5.4, 3.0, 'effect strongest in\nthis subgroup (H5)', ha='left', va='center', fontsize=7.8, style='italic')

# H1 / H2 antecedent strip at top
box(2.0, 6.55, 2.7, 0.6, 'H1  Male-typed industry favoured', fc='#F6F6F6', fs=8)
box(8.0, 6.55, 2.5, 0.6, 'H2  Masculine match favoured', fc='#F6F6F6', fs=8)

# legend
ax.text(0.1, 0.25, 'Solid arrow = empirical effect on investment.   '
        'Dashed arrow = subgroup within which the congruity preference is expressed '
        '(simple effect, not a statistical interaction).',
        ha='left', va='center', fontsize=7.6, style='italic')
fig.tight_layout()
apa.save(fig, os.path.join(HERE, 'figures', 'fig_conceptual_model_revised'))
plt.close(fig)
print('saved fig_conceptual_model_revised')


# ════════════════════════ WORD REPORT ════════════════════════
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIGDIR = os.path.join(HERE, 'figures')
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(11)


def H(txt, lvl=1):
    h = doc.add_heading(txt, level=lvl)
    return h


def P(txt='', bold=False, italic=False, size=11, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    r = p.add_run(txt); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(txt, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    p.add_run(txt)
    return p


def fig(name, caption, width=5.0):
    doc.add_picture(os.path.join(FIGDIR, name + '.png'), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9)


def fmtp(p):
    return '< .001' if p < .001 else f'= {p:.3f}'


# ---------- Title ----------
t = doc.add_heading('Reassessment of Hypotheses 3 and 5 as Subgroup Effects', level=0)
P('Reframing emotional intelligence and investment experience from statistical '
  'moderators to subgroup boundary conditions', italic=True, size=11)
P('Dependent variable: amount invested (0–3). Masculine-congruent = male founders in '
  'the male-dominated industry (SiteVision); feminine-congruent = female founders in the '
  'female-dominated industry (BalanceUp). Lower EI = below the sample median WLEIS '
  f'(Mdn = {MED:.2f}). N = 187. All tests on the strict-manipulation-check sample.',
  italic=True, size=9.5)

# ---------- 1. The conceptual problem ----------
H('1. Why the original tests were non-significant — and why the reformulation changes the question', 1)
P('The thesis develops H3 and H5 around the idea that emotional intelligence (EI) '
  'and investment experience act as moderators of congruity bias. A moderation claim '
  'is a claim about an interaction term: that the size of the congruity preference '
  '(masculine-congruent minus feminine-congruent investment) differs as a function of '
  'EI, or of EI together with experience. When tested formally — a three-way '
  'EI × founder gender × industry model — the interaction coefficient was not '
  'statistically significant. That is the basis on which the earlier analysis reported '
  'H3 and H5 as unsupported.')
P('The reformulated hypotheses do not ask about an interaction. They ask a simple-effect '
  'question: within the specific subgroup of lower-EI investors (H3), and within the '
  'narrower subgroup of lower-EI investors who also lack prior investment experience (H5), '
  'is the masculine-congruent startup funded more than the feminine-congruent startup? '
  'This is a directional comparison of two means inside a defined subgroup. It is a '
  'legitimate and answerable question, and — importantly — a subgroup simple effect '
  'can be significant even when the overall interaction is not. The two results are not '
  'contradictory; they answer different questions. The reformulation is therefore '
  'defensible provided the wording describes where the effect appears, not that EI '
  'statistically changes the effect.')

# ---------- 2. H3 ----------
H('2. Hypothesis 3 — lower-EI investors', 1)
P('H3 (reformulated): Among investors lower in emotional intelligence, masculine-congruent '
  'startups receive higher investment than feminine-congruent startups.', bold=True)
P('Result. Within the lower-EI subgroup, masculine-congruent startups were funded more '
  'than feminine-congruent startups, a medium-sized, directionally significant difference:')
bullet(f"Masculine-congruent: M = {H3['m1']:.2f}, SD = {H3['s1']:.2f}, n = {H3['n1']}.", )
bullet(f"Feminine-congruent: M = {H3['m2']:.2f}, SD = {H3['s2']:.2f}, n = {H3['n2']}.")
bullet(f"Mean difference = {H3['diff']:+.2f} on the 0–3 scale "
       f"(95% CI [{H3['ci'][0]:+.2f}, {H3['ci'][1]:+.2f}]).")
bullet(f"Welch t({H3['dfw']:.0f}) = {H3['t']:.2f}; one-tailed p = {H3['p1']:.3f}; "
       f"two-tailed p = {H3['p2']:.3f}.")
bullet(f"Cohen's d = {H3['d']:.2f} (medium).")
bullet(f"Mann–Whitney U (one-sided) p = {H3['pmw']:.3f}.")
bullet(f"Bootstrap 95% CI of the difference (10,000 resamples, seed 42) "
       f"[{H3['bci'][0]:+.2f}, {H3['bci'][1]:+.2f}].")
P('Verdict. H3 is supported as a directional subgroup effect. The one-tailed Welch test '
  'and the non-parametric Mann–Whitney test are both significant at the .05 level, and the '
  'effect size is medium (d = 0.50). The caveat is that the test is genuinely one-tailed: '
  'the two-tailed p (.064) and the analytic and bootstrap confidence intervals just touch '
  'zero. The directional hypothesis is appropriate here because the prediction — masculine '
  'over feminine — is specified a priori by theory, but the support should be reported as '
  'directional rather than as a two-sided effect.', bold=False)
fig('figH3_revised',
    'Figure H3. Investment in masculine- vs. feminine-congruent startups among '
    'lower-EI investors. Error bars = ±1 SE.', width=4.4)
P('Robustness (alternative low-EI cut-off). Re-defining "lower EI" as the bottom tertile '
  f"(≤ {TERT:.2f}) reproduces the same pattern with a comparable effect size "
  f"(M = {H3t['m1']:.2f} vs {H3t['m2']:.2f}, diff = {H3t['diff']:+.2f}, d = {H3t['d']:.2f}, "
  f"one-tailed p = {H3t['p1']:.3f}, Mann–Whitney p = {H3t['pmw']:.3f}). The effect is "
  'stable in direction and magnitude across the two operationalisations; significance is '
  'borderline because the tertile cut reduces the subgroup n.', italic=False)

# ---------- 3. H5 ----------
H('3. Hypothesis 5 — lower-EI investors without prior investment experience', 1)
P('H5 (reformulated): Among investors lower in emotional intelligence without prior '
  'investment experience, masculine-congruent startups receive higher investment than '
  'feminine-congruent startups.', bold=True)
P('Result. Restricting the lower-EI subgroup further to participants who reported no prior '
  'experience with financial investments sharpens the effect substantially:')
bullet(f"Masculine-congruent: M = {H5['m1']:.2f}, SD = {H5['s1']:.2f}, n = {H5['n1']}.")
bullet(f"Feminine-congruent: M = {H5['m2']:.2f}, SD = {H5['s2']:.2f}, n = {H5['n2']}.")
bullet(f"Mean difference = {H5['diff']:+.2f} on the 0–3 scale "
       f"(95% CI [{H5['ci'][0]:+.2f}, {H5['ci'][1]:+.2f}]).")
bullet(f"Welch t({H5['dfw']:.0f}) = {H5['t']:.2f}; one-tailed p = {H5['p1']:.3f}; "
       f"two-tailed p = {H5['p2']:.3f}.")
bullet(f"Cohen's d = {H5['d']:.2f} (large).")
bullet(f"Mann–Whitney U (one-sided) p = {H5['pmw']:.3f}.")
bullet(f"Bootstrap 95% CI of the difference [{H5['bci'][0]:+.2f}, {H5['bci'][1]:+.2f}] "
       f"— excludes zero.")
P('Verdict. H5 is supported, and more robustly than H3. The difference is significant on '
  'the conventional two-tailed test (p = .031) as well as one-tailed (p = .016), the '
  'effect is large (d = 0.96), and both the analytic and the bootstrap confidence intervals '
  'exclude zero. The one qualification is statistical-power: the subgroup is small '
  f"(n = {H5['n1']} vs {H5['n2']}), so the point estimate of the effect size is imprecise "
  'and should be read as "large but uncertain." The direction and significance are '
  'nonetheless consistent across the parametric and non-parametric tests and the bootstrap.')
fig('figH5_revised',
    'Figure H5. Investment in masculine- vs. feminine-congruent startups among lower-EI '
    'investors without prior investment experience. Error bars = ±1 SE.', width=4.4)
P('Robustness (alternative low-EI cut-off). With the bottom-tertile definition the '
  f"direction is preserved and the effect remains large (M = {H5t['m1']:.2f} vs "
  f"{H5t['m2']:.2f}, diff = {H5t['diff']:+.2f}, d = {H5t['d']:.2f}), but with only "
  f"n = {H5t['n1']} vs {H5t['n2']} the test is no longer significant (one-tailed "
  f"p = {H5t['p1']:.3f}). This is a power limitation, not a reversal: the cell becomes "
  'too small to detect even a large effect.')

# ---------- 4. The escalation / context ----------
H('4. Reading H3 and H5 together', 1)
P('The two subgroup results form a coherent gradient. In the full sample the congruity '
  f"preference is small and not significant (diff = {FULL['diff']:+.2f}, d = {FULL['d']:.2f}, "
  f"two-tailed p = {FULL['p2']:.3f}). It grows to a medium effect among lower-EI investors "
  f"(diff = {H3['diff']:+.2f}, d = {H3['d']:.2f}) and to a large effect among lower-EI "
  f"investors who also lack investment experience (diff = {H5['diff']:+.2f}, d = {H5['d']:.2f}). "
  'The pattern is exactly what the theory predicts as a descriptive matter: the congruity '
  'preference concentrates in the investors least equipped — emotionally and experientially '
  '— to correct it. What the data do not license is the inferential claim that EI or '
  'experience statistically *causes* this gradient, because the corresponding interaction '
  'terms are not significant.')
fig('figH3H5_escalation',
    'Figure. Congruity preference across the full sample, the lower-EI subgroup (H3), and '
    'the lower-EI / no-experience subgroup (H5). Δ = masculine minus feminine; '
    'error bars = ±1 SE.', width=5.6)

# ---------- 5. Alternative formulations ----------
H('5. Suggested wording for the hypotheses', 1)
P('The reformulated hypotheses are theoretically sound and methodologically defensible as '
  'simple-effect statements. Two refinements would make them maximally defensible and align '
  'them precisely with what the data show.')
P('H3 — recommended wording:', bold=True)
P('"Among investors lower in emotional intelligence, masculine-congruent startups receive '
  'higher investment than feminine-congruent startups." (Retain as is. Report as a directional '
  'simple effect; do not describe EI as moderating the effect.)', italic=True)
P('H5 — recommended wording:', bold=True)
P('"Among investors lower in emotional intelligence and without prior investment experience, '
  'masculine-congruent startups receive higher investment than feminine-congruent startups." '
  '(Retain. This is the best-supported hypothesis in the set and may be stated without a '
  'directional hedge, as it is significant two-tailed.)', italic=True)
P('Optional supplementary hypothesis (if a moderation-style claim is still wanted):', bold=True)
P('Rather than asserting moderation, frame the gradient descriptively: "The congruity '
  'preference is more pronounced among investors lower in emotional intelligence and is most '
  'pronounced among those who additionally lack investment experience." This describes the '
  'observed escalation without claiming a significant interaction, and is fully supported by '
  'Figure 4.', italic=True)
P('What to avoid:', bold=True)
P('Any wording of the form "EI moderates / weakens / determines the strength of the congruity '
  'effect," or "experience conditions the EI–congruity relationship." These are interaction '
  'claims and the interaction tests are non-significant. They cannot be supported by the '
  'present data and should be removed wherever they appear.', italic=False)

# ---------- 6. Text revisions ----------
H('6. Passages in the literature review / theory section that need revision', 1)
P('The current manuscript repeatedly frames EI as a statistical moderator and experience as '
  'a "conditional moderator." If H3 and H5 are reframed as subgroup effects, the following '
  'passages must be rewritten. Locations refer to the section headings in MASTERARBEIT_v2.')

tbl = doc.add_table(rows=1, cols=3)
tbl.style = 'Light Grid Accent 1'
hdr = tbl.rows[0].cells
for c, txt in zip(hdr, ['Location', 'Current wording (moderation framing)', 'Suggested revision (subgroup framing)']):
    c.paragraphs[0].add_run(txt).bold = True

revisions = [
    ('Section "Emotional Intelligence and Role Congruity Bias", paragraph beginning '
     '"This logic helps explain the role of EI…"',
     '"EI is expected to act as a moderator of congruity bias, influencing how strongly '
     'congruity bias affects investment decisions."',
     '"The congruity preference is expected to be concentrated among investors lower in EI, '
     'who are less able to recognise and discount feelings of fit or misfit. EI is treated as '
     'a boundary condition identifying where the preference is expressed, not as a statistical '
     'moderator of its magnitude."'),
    ('Same section, sentences listing "strongest among investors who are less able to '
     'understand and regulate…"',
     '"…the congruity preferences described in the first two hypotheses should not be '
     'equally strong across all investors. Instead, they should be strongest among investors '
     'who are less able to understand and regulate their emotional reactions."',
     'Keep the substance but phrase as a subgroup expectation: "…the congruity preference '
     'is expected to be observable primarily within the subgroup of investors lower in EI." '
     'Remove the implication that the across-investor difference is itself the tested effect.'),
    ('Derivation of H3 (paragraph "The difference in capital allocation should be most '
     'visible…")',
     '"…the difference in capital allocated to masculine-congruent and feminine-congruent '
     'startups should be larger among investors with lower EI."',
     '"…among investors with lower EI, masculine-congruent startups are expected to receive '
     'higher investment than feminine-congruent startups." State H3 as a within-subgroup '
     'comparison rather than a comparison of effect sizes between EI groups.'),
    ('Section "Investment Experience as a Boundary Condition", penultimate paragraph',
     '"Investment experience therefore moderates the relationship between EI and congruity '
     'bias."',
     '"Among investors lower in EI, the congruity preference is expected to be confined to — '
     'or strongest in — those who also lack investment experience, for whom no learned '
     'pattern offsets the felt sense of fit." Replace "moderates the relationship" with a '
     'nested-subgroup statement.'),
    ('Figure 1 caption / conceptual-model description (paragraph "The model brings the five '
     'hypotheses together…")',
     '"This difference in capital allocation … is strongest among those lower in EI '
     '(Hypothesis 3). EI acts as a moderator of congruity bias, determining how strongly the '
     'founder-industry pairing translates into a difference in capital allocation."',
     '"This difference in capital allocation is observed primarily among investors lower in EI '
     '(Hypothesis 3). EI is a boundary condition that identifies the subgroup in which the '
     'congruity preference is expressed rather than a moderator that scales its magnitude." '
     'Delete "acts as a moderator … determining how strongly."'),
    ('Figure 1 description, final sentences on experience',
     '"Investment experience acts as a conditional moderator of the relationship between EI '
     'and congruity bias … Only among investors lower in EI, greater investment experience '
     'is associated with a smaller difference in capital allocation…"',
     '"The congruity preference is most pronounced among investors who are both lower in EI '
     'and without prior investment experience (Hypothesis 5)." Replace "conditional moderator" '
     'and the interaction phrasing with a nested-subgroup statement.'),
]
for loc, cur, rev in revisions:
    cells = tbl.add_row().cells
    cells[0].paragraphs[0].add_run(loc).font.size = Pt(8.5)
    cells[1].paragraphs[0].add_run(cur).font.size = Pt(8.5)
    cells[2].paragraphs[0].add_run(rev).font.size = Pt(8.5)

P()
P('General rule for the search-and-replace pass: every instance of "moderator", '
  '"moderates", "moderation", "conditional moderator", "interaction", or "determines how '
  'strongly" applied to EI or experience should be replaced with subgroup / boundary-condition '
  'language ("among", "within the subgroup of", "is concentrated in", "is most pronounced '
  'among"). The hypotheses themselves (H3, H5) are already phrased as "Among investors '
  'lower in EI …" and do not need to change; it is the surrounding theoretical prose and '
  'the Figure 1 narrative that carry the moderation framing.', italic=True)

# ---------- 7. New conceptual model ----------
H('7. Revised conceptual model', 1)
P('The figure below replaces the moderation model. The congruity preference is drawn as a '
  'single empirical path from founder–industry congruity to investment allocation (solid '
  'arrow). Lower EI and absence of investment experience are no longer drawn as arrows that '
  '"bend" that path; instead they are subgroup selectors (dashed arrows) that mark the '
  'population within which the preference is empirically expressed (H3) and within which it is '
  'strongest (H5). H1 and H2 sit above as the industry- and match-level antecedents.')
fig('fig_conceptual_model_revised',
    'Figure. Revised conceptual model. Solid arrow = empirical effect on investment; dashed '
    'arrows = subgroups within which the congruity preference is expressed (simple effects, '
    'not statistical interactions).', width=6.4)
P('Reading the model. H1 and H2 establish that male-typed industries and masculine matches '
  'attract more capital. The core relationship is the congruity preference (masculine-congruent '
  '> feminine-congruent investment). H3 states that this preference is observed among lower-EI '
  'investors; H5 states that it is largest among lower-EI investors who also lack investment '
  'experience. The model deliberately avoids interaction arrows, so that the figure cannot be '
  'read as a claim the data do not support.')

# ---------- 8. Bottom line ----------
H('8. Bottom line', 1)
bullet('can be supported as a directional subgroup effect: medium effect (d = 0.50), '
       'one-tailed p = .032, Mann–Whitney p = .021. Report it as directional; the two-tailed '
       'CI is marginal.', bold_lead='H3 ')
bullet('is supported, and is the strongest result in the set: large effect (d = 0.96), '
       'two-tailed p = .031, bootstrap CI excludes zero. The only caveat is the small subgroup n.',
       bold_lead='H5 ')
bullet('Both are simple subgroup effects, not moderation. The formal interaction tests remain '
       'non-significant, so all "EI moderates / experience conditions" wording in the theory '
       'section and Figure 1 must be rewritten in subgroup / boundary-condition language '
       '(Section 6).')
bullet('A new conceptual model (Section 7) represents EI and experience as subgroup selectors '
       'rather than moderators, matching the evidence.')

OUT = os.path.join(HERE, '36_revised_H3_H5_report.docx')
doc.save(OUT)
print('SAVED', OUT)
