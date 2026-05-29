#!/usr/bin/env python3
"""Build a Word memo: Literature Review vs. Experimental Results comparison.
Output: analysis/26_litreview_comparison_memo.docx
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'; style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)

def body(t, italic=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.italic = italic
    p.paragraph_format.line_spacing = 1.5
    return p

def bullet(t):
    p = doc.add_paragraph(t, style='List Bullet')
    p.paragraph_format.line_spacing = 1.3
    return p

def simple_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Light List Accent 1'
    for i, h in enumerate(headers):
        run = t.rows[0].cells[i].paragraphs[0].add_run(h); run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    return t

doc.add_heading('Literature Review vs. Experimental Results: Comparison and Required Additions', level=1)
body(
    'This memo compares the claims advanced in the literature review with the '
    'findings of the present experiment (N = 187) and identifies the contradicting '
    'and approving literature that should be incorporated. The central observation '
    'is that the review is structured to predict effects that the data largely did '
    'not produce; the most important additions are therefore boundary-condition '
    'framing and a small set of contradicting citations, not only confirmatory ones.'
)

# ── Confirmations ──
doc.add_heading('1. Where the Results Confirm the Review', level=2)
simple_table(
    ['Review claim', 'Experimental result'],
    [['Industry gender-typing is a key boundary condition; male-typed industries '
      '(IT/ConTech) are advantaged (Sperber & Linder; Tonoyan & Strohmeyer).',
      'R1: SiteVision (male-typed ConTech) >> BalanceUp (female-typed Wellness), '
      'paired t = 5.08, p < .001 — the strongest and most robust effect.'],
     ['Women in female-typed or gender-neutral industries are not penalised, and '
      'may even be overvalued (Tonoyan & Strohmeyer 2021).',
      'No within-industry penalty for female founders; no raw male-vs-female gap '
      '(p = .257), with female-founded ventures rated slightly higher.'],
     ['EI is linked to risk perception and the handling of affect in financial '
      'decisions (Yip & Cote; Bucciol et al.).',
      'Perceived success probability and risk were the operative covariates in the '
      'robustness checks; the congruity effect tracked perceived viability.']])

# ── Contradictions ──
doc.add_heading('2. Where the Results Contradict the Review (priority issues)', level=2)
body('2.1 The gender funding gap as a main effect.', italic=False)
body(
    'The review opens by asserting a robust funding gap disadvantaging women. The '
    'experiment found no raw male-versus-female founder gap (p = .257; if anything '
    'reversed). The funding difference resides entirely in the founder-by-industry '
    'congruity cells, not in founder gender per se. The review needs a paragraph '
    'conceding that experimental and lab evidence is mixed and that the gap is '
    'conditional on industry congruity rather than a universal main effect.'
)
body('2.2 The warmth-deficit premise (SCM and the H9 mediation chain).', italic=False)
body(
    'The review leans on Eckes and Cuddy et al.: female founders elicit low warmth '
    '(the envious stereotype), and warmth-driven emotions mediate behaviour. In the '
    'data, female-founded teams were rated equal-to-slightly-higher on warmth (ns), '
    'and warmth did not mediate investment (bootstrap 95% CI included zero). The '
    'operative mediator was perceived success/viability — closer to a competence '
    'pathway. This is the largest theoretical exposure: the SCM-to-warmth chain '
    'should be reframed toward competence-based mechanisms, and warmth mediation '
    'flagged as an exploratory rather than strongly predicted test.'
)
body('2.3 The Self-Emotion-Appraisal (SEA) thesis.', italic=False)
body(
    'The review argues at length that SEA is the dimension driving EI’s debiasing '
    'role. In the data, the SEA x congruity interaction was b = +0.05, p = .77 — '
    'not even directionally correct, and no EI dimension moderated significantly. '
    'The "Specific Role of Self-Emotion Appraisal" section currently overcommits '
    'and should be softened to a hypothesis, supported by literature noting that '
    'EI-dimension effects on decision-making are weak and inconsistent.'
)
body('2.4 The "women are more emotional" stereotype.', italic=False)
body(
    'The review presents the stereotype that women are more emotional and less able '
    'to regulate emotion. In this sample, women scored significantly higher on EI '
    '(total, SEA, OEA; d up to 0.66) — an ironic empirical counterpoint worth '
    'noting in the discussion: the group stereotyped as "too emotional" showed '
    'higher measured emotional ability.'
)

# ── Papers to add ──
doc.add_heading('3. Literature to Add', level=2)
body('Contradicting / nuancing (to pre-empt reviewers):', italic=False)
bullet('The Science Advances (2020) industry-fit study — investors penalise female '
       'founders for lack of industry fit, with no fit penalty for male founders. '
       'Engage it directly: our data show the industry effect but not the '
       'female-specific fit penalty. (science.org/doi/10.1126/sciadv.abd7664)')
bullet('Crowdfunding warmth/competence work showing female-led campaigns raise '
       'comparable total funding and that competence — not warmth — drives women’s '
       'success (Svetek 2023; Anglin et al. crowdfunding warmth/competence). '
       'Directly relevant to the null warmth-mediation result.')
bullet('2024 ability-EI decision-making review noting dimension-specific EI effects '
       'are weak/inconsistent in non-social, emotionally charged tasks — cover for '
       'the null SEA result. (PMC12070655)')
body('Confirming:', italic=False)
bullet('Tonoyan & Strohmeyer (2021) — lean on it harder; it predicts the observed '
       '"no penalty in the female-typed industry" pattern.')
bullet('EI–risk-perception serial-mediation work, to connect the risk/success '
       'findings to the EI literature. (PMC9204191)')

# ── Edits ──
doc.add_heading('4. Concrete Edits to Make', level=2)
bullet('Add a "mixed evidence / boundary conditions" paragraph in The Gender '
       'Funding Gap: the gap is not universal, experimental results are mixed, and '
       'it is conditional on industry congruity.')
bullet('Reframe the SCM section to foreground competence/viability alongside '
       'warmth, and mark the warmth-mediation (H9) as exploratory.')
bullet('Demote the SEA subsection from "the mechanism" to "a candidate mechanism," '
       'adding the EI-dimension-inconsistency citation.')
bullet('Add a sentence acknowledging that the lab sample (students, hypothetical '
       'small allocations, no real stakes) and the within-person two-pitch design '
       'are boundary conditions that may attenuate the gender gap — this honestly '
       'explains the null results.')

# ── References ──
doc.add_heading('Key Sources', level=2)
for s in [
    'Evidence that investors penalize female founders for lack of industry fit. '
    'Science Advances (2020). https://www.science.org/doi/10.1126/sciadv.abd7664',
    'Svetek, M. (2023). Perceived competence and cooperativeness in early-stage '
    'financing. ET&P. https://journals.sagepub.com/doi/10.1177/10422587221127000',
    'Perceived warmth and competence in crowdfunding: which matters more and for '
    'whom? JBV Insights (2022). '
    'https://www.sciencedirect.com/science/article/abs/pii/S2352673421000664',
    'The roles of ability emotional intelligence in predicting affective '
    'decision-making (2024). https://pmc.ncbi.nlm.nih.gov/articles/PMC12070655/',
    'Optimism bias and risk perception mediating EI and decision-making. '
    'https://www.ncbi.nlm.nih.gov/pmc/articles/PMC9204191/',
    'Gender homophily and the stigma of incompetence in entrepreneurial finance. '
    'Organization Science. https://pubsonline.informs.org/doi/10.1287/orsc.2022.1594',
]:
    p = doc.add_paragraph(s); p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(4)

OUT = os.path.join(HERE, '26_litreview_comparison_memo.docx')
doc.save(OUT)
print('Saved', OUT)
