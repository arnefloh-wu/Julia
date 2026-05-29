#!/usr/bin/env python3
"""Build a Word draft of the Sample Description section.

Computes all demographic figures from the strict cleaned dataset (N=187),
writes an English narrative, embeds Figure 1 (gender bar chart) and the
demographic tables. Output: analysis/02_sample_description_draft.docx
"""
import csv, statistics as st, sys, os
from collections import Counter
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import labels as L

# ---------- data ----------
SRC = os.path.join(os.path.dirname(HERE), 'MasterThesis_cleaned.csv')
with open(SRC, newline='') as f:
    rows = list(csv.reader(f))
names, data = rows[0], rows[3:]
N = len(data)
ci = {q: names.index(q) for q in
      ['Q16', 'Q17', 'Q18', 'Q19', 'Q20', 'Q21', 'Q22', 'Q23', 'Q24']}

def vget(r, q): return r[ci[q]].strip()

def freq(q, labelmap=None, order=None):
    """Return (rows, valid_n, missing) where rows = [(label, n, pct_valid)]."""
    c = Counter(vget(r, q) for r in data if vget(r, q) != '')
    missing = sum(1 for r in data if vget(r, q) == '')
    base = sum(c.values())
    keys = order if order else [k for k, _ in c.most_common()]
    out = []
    for k in keys:
        if k in c:
            lab = labelmap.get(k, k) if labelmap else k
            out.append((lab, c[k], c[k] / base * 100))
    return out, base, missing

# age
age_fix = {'20 Jahre': '20'}
ages, agrp = [], Counter()
for r in data:
    v = age_fix.get(vget(r, 'Q17'), vget(r, 'Q17'))
    try:
        a = int(float(v)); ages.append(a)
        g = ('<20' if a < 20 else '20-24' if a < 25 else '25-29' if a < 30
             else '30-39' if a < 40 else '40+')
        agrp[g] += 1
    except ValueError:
        if '40' in v:
            agrp['40+'] += 1
age_M, age_SD = st.mean(ages), st.stdev(ages)
age_md, age_min, age_max = int(st.median(ages)), min(ages), max(ages)

gender, _, _ = freq('Q16', L.GENDER, L.GENDER_ORDER)
edu, edu_n, edu_miss = freq('Q19', L.EDUCATION, L.EDUCATION_ORDER)
deg, deg_n, deg_miss = freq('Q20', L.DEGREE, L.DEGREE_ORDER)
field, _, _ = freq('Q21', L.FIELD)
wu, wu_n, wu_miss = freq('Q22', L.YESNO, ['Ja', 'Nein'])
inv, _, _ = freq('Q23', L.YESNO, ['Ja', 'Nein'])
found, _, _ = freq('Q24', L.YESNO, ['Ja', 'Nein'])
nat = Counter()
for r in data:
    v = vget(r, 'Q18')
    if v in L.NATIONALITY:
        nat[L.NATIONALITY[v]] += 1
nat_rows = [(k, n, n / N * 100) for k, n in nat.most_common()]

# convenience lookups for prose
gd = {lab: (n, p) for lab, n, p in gender}
ed = {lab: (n, p) for lab, n, p in edu}
dd = {lab: (n, p) for lab, n, p in deg}

# ---------- document ----------
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)

def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing = 1.5
    return p

doc.add_heading('Sample Description', level=1)

body(
    f'A total of {N} participants were included in the final sample. Of the '
    f'230 recorded responses, 43 were excluded because they failed one of two '
    f'manipulation checks: the correct recall of the start-ups’ industries '
    f'(Q7) and the correct recall of the founding teams’ gender composition '
    f'(Q8). Participants who indicated that they could not remember this '
    f'information were likewise removed. Unless stated otherwise, the '
    f'percentages reported below are based on the full sample of N = {N}; for '
    f'three variables a small number of participants did not provide an answer, '
    f'in which case valid percentages and the respective sample size are '
    f'reported.'
)

# --- Gender + Figure 1 ---
body(
    f'The sample comprised a slight majority of female participants. As shown '
    f'in Figure 1, {gd["Female"][0]} participants identified as female '
    f'({gd["Female"][1]:.1f}%) and {gd["Male"][0]} as male '
    f'({gd["Male"][1]:.1f}%), while {gd["Non-binary"][0]} participants '
    f'({gd["Non-binary"][1]:.1f}%) identified as non-binary and '
    f'{gd["Prefer not to say"][0]} ({gd["Prefer not to say"][1]:.1f}%) '
    f'preferred not to disclose their gender.'
)

cap = doc.add_paragraph()
cap.paragraph_format.space_before = Pt(6)
r = cap.add_run('Figure 1'); r.bold = True
cap.add_run('\n')
r = cap.add_run('Gender Distribution of the Sample'); r.italic = True
img = doc.add_paragraph()
img.alignment = WD_ALIGN_PARAGRAPH.CENTER
img.add_run().add_picture(os.path.join(HERE, 'figures', 'fig01_gender.png'),
                          width=Inches(5.8))
note = doc.add_paragraph()
nr = note.add_run('Note. '); nr.italic = True
note.add_run(f'N = {N}.')
note.paragraph_format.space_after = Pt(12)

# --- Age ---
body(
    f'Participants were on average {age_M:.1f} years old (SD = {age_SD:.1f}, '
    f'Mdn = {age_md}), ranging from {age_min} to {age_max} years. The majority '
    f'fell into the 20–24 age bracket ({agrp["20-24"]/N*100:.1f}%), '
    f'followed by the 25–29 group ({agrp["25-29"]/N*100:.1f}%), reflecting '
    f'the predominantly student composition of the sample.'
)

# --- Nationality ---
top = nat.most_common()
body(
    f'With respect to nationality, the sample was dominated by German '
    f'({nat["German"]/N*100:.1f}%) and Austrian ({nat["Austrian"]/N*100:.1f}%) '
    f'participants, who together accounted for roughly two thirds of '
    f'respondents. The remaining third comprised a broad range of '
    f'nationalities, most notably Ukrainian ({nat["Ukrainian"]/N*100:.1f}%), '
    f'Bulgarian ({nat["Bulgarian"]/N*100:.1f}%), and Russian '
    f'({nat["Russian"]/N*100:.1f}%). The full distribution is reported in '
    f'Table 2.'
)

# --- Education / field / WU ---
field_pct = {l: p for l, n, p in field}
wu_pct = {l: p for l, n, p in wu}
body(
    f'In terms of education, most participants reported a higher-education '
    f'entrance qualification (Matura/Abitur) as their highest completed '
    f'degree ({ed[L.EDUCATION["Matura / Abitur / Hochschulreife"]][1]:.1f}%), '
    f'followed by a Bachelor’s degree ({ed[L.EDUCATION["Bachelor"]][1]:.1f}%; '
    f'valid n = {edu_n}). Consistent with this, the majority were currently '
    f'pursuing a Bachelor’s ({dd[L.DEGREE["Bachelor"]][1]:.1f}%) or a '
    f'Master’s degree ({dd[L.DEGREE["Master"]][1]:.1f}%). The most common field '
    f'of study was Business Administration/Management '
    f'({field_pct["Business Administration/Management"]:.1f}%), '
    f'and {wu_pct["Yes"]:.1f}% of participants were '
    f'enrolled at the Vienna University of Economics and Business (WU; '
    f'valid n = {wu_n}).'
)

# --- Experience ---
inv_pct = {l: p for l, n, p in inv}
found_pct = {l: p for l, n, p in found}
body(
    f'Finally, {inv_pct["Yes"]:.1f}% of participants '
    f'reported prior experience with financial investments, whereas only '
    f'{found_pct["Yes"]:.1f}% had previously founded or '
    f'actively participated in founding a company. A complete overview of the '
    f'sample’s demographic characteristics is provided in Table 1.'
)

# ---------- Table 1 ----------
def caption(num, title):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    r = p.add_run(num); r.bold = True
    p.add_run('\n'); r = p.add_run(title); r.italic = True

def stat_table(blocks, note_text=None):
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Light List Accent 1'
    hdr = t.rows[0].cells
    for i, h in enumerate(['Characteristic', 'n', '%']):
        run = hdr[i].paragraphs[0].add_run(h); run.bold = True
    for title, rws in blocks:
        cells = t.add_row().cells
        rr = cells[0].paragraphs[0].add_run(title); rr.bold = True
        for lab, n, p in rws:
            c = t.add_row().cells
            c[0].text = '   ' + lab
            c[1].text = str(n)
            c[2].text = f'{p:.1f}'
    if note_text:
        np = doc.add_paragraph(); nr = np.add_run('Note. '); nr.italic = True
        np.add_run(note_text)

caption('Table 1', 'Demographic Characteristics of the Sample (N = 187)')
age_rows = [(g, agrp[g], agrp[g] / N * 100)
            for g in ['<20', '20-24', '25-29', '30-39', '40+'] if agrp[g]]
stat_table([
    ('Gender', gender),
    ('Age group', age_rows),
    (f'Highest completed education (valid n = {edu_n})', edu),
    (f'Degree currently pursued (valid n = {deg_n})', deg),
    ('Field of study', field),
    (f'Enrolled at WU Vienna (valid n = {wu_n})', wu),
    ('Experience with financial investments', inv),
    ('Entrepreneurial experience', found),
],
    note_text=(f'Age: M = {age_M:.1f}, SD = {age_SD:.1f}, range {age_min}–{age_max}. '
               f'Percentages are based on N = {N} unless a valid n is indicated; '
               f'the {edu_miss} participants missing on education, degree, and WU '
               f'enrolment correspond to an earlier survey version.'))

# ---------- Table 2: nationality ----------
caption('Table 2', 'Nationality of Participants (N = 187)')
t = doc.add_table(rows=1, cols=3)
t.style = 'Light List Accent 1'
for i, h in enumerate(['Nationality', 'n', '%']):
    run = t.rows[0].cells[i].paragraphs[0].add_run(h); run.bold = True
for lab, n, p in nat_rows:
    c = t.add_row().cells
    c[0].text = lab; c[1].text = str(n); c[2].text = f'{p:.1f}'

OUT = os.path.join(HERE, '02_sample_description_draft.docx')
doc.save(OUT)
print('Saved', OUT)
print(f'N={N} | gender={ {l:n for l,n,_ in gender} } | age M={age_M:.1f} SD={age_SD:.1f}')
